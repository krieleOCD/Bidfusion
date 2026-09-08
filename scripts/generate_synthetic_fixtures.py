#!/usr/bin/env python3
"""Generate the synthetic RFP fixture packages under ``fixtures/synthetic/``.

All content is invented test data. The PDF builder writes byte-deterministic
minimal PDFs, so regenerating the PDF fixtures must reproduce the committed
bytes exactly. XLSX output is content-deterministic (fixed document
properties), but ZIP entry timestamps vary per run, so XLSX fixtures are
compared by content, never by bytes.

Usage: ``python scripts/generate_synthetic_fixtures.py [--output DIR]``
"""

from __future__ import annotations

import argparse
import zipfile
from collections.abc import Sequence
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = ROOT / "fixtures" / "synthetic"

INVITATION_PAGE_1 = (
    "Synthetic RFP BF-2026-DEMO - Invitation to Tender",
    "Customer: Example Testing Corp (synthetic test data)",
    "This request for proposal is fully synthetic and contains no real data.",
    "Submission deadline: 2026-09-30 17:00",
    "Q&A deadline: 2026-09-11 12:00",
    "Offers must be submitted via the Example procurement portal.",
    "Proposals must be submitted in PDF format.",
    "The offer letter must be signed by an authorized representative.",
)
INVITATION_PAGE_2 = (
    "Requirements",
    "REQ-P1: The supplier MUST provide a 24x7 security operations center.",
    "REQ-P2: The supplier MUST encrypt customer data at rest and in transit.",
    "REQ-P3: The service SHOULD include monthly vulnerability reports.",
    "REQ-P4: The supplier MAY offer an optional on-site workshop.",
    "REQ-P5: The solution handles peak load appropriately.",
    "Contract start: 2027-01-01",
)

_FIXED_XLSX_TIMESTAMP = datetime(2026, 1, 1, 12, 0, 0)


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _page_content_stream(lines: Sequence[str]) -> bytes:
    operations = []
    y = 720
    for line in lines:
        operations.append(f"BT /F1 11 Tf 72 {y} Td ({_pdf_escape(line)}) Tj ET")
        y -= 18
    return "\n".join(operations).encode("ascii")


def build_pdf(pages: Sequence[Sequence[str]], *, encrypt_marker: bool = False) -> bytes:
    """Build a deterministic, minimal text PDF from per-page line lists."""
    if not pages:
        msg = "a PDF fixture requires at least one page"
        raise ValueError(msg)

    page_count = len(pages)
    font_object = 3 + 2 * page_count
    encrypt_object = font_object + 1 if encrypt_marker else None
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(page_count))

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode("ascii"),
    ]
    for index, lines in enumerate(pages):
        page_object = 3 + 2 * index
        contents_object = page_object + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_object} 0 R >> >> "
                f"/Contents {contents_object} 0 R >>"
            ).encode("ascii")
        )
        stream = _page_content_stream(lines)
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    if encrypt_object is not None:
        objects.append(b"<< /Filter /Standard /V 1 /R 2 >>")

    buffer = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(buffer))
        buffer += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"

    xref_offset = len(buffer)
    buffer += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    buffer += b"0000000000 65535 f \n"
    for offset in offsets:
        buffer += f"{offset:010d} 00000 n \n".encode("ascii")

    trailer = f"<< /Size {len(objects) + 1} /Root 1 0 R"
    if encrypt_object is not None:
        trailer += f" /Encrypt {encrypt_object} 0 R"
    trailer += " >>"
    buffer += (
        b"trailer\n"
        + trailer.encode("ascii")
        + f"\nstartxref\n{xref_offset}\n".encode("ascii")
        + b"%%EOF\n"
    )
    return bytes(buffer)


def build_invitation_pdf() -> bytes:
    """The clean two-page synthetic RFP invitation."""
    return build_pdf([INVITATION_PAGE_1, INVITATION_PAGE_2])


def build_blank_scan_pdf() -> bytes:
    """A structurally valid PDF without any extractable text."""
    return build_pdf([[]])


def build_encrypted_marker_pdf() -> bytes:
    """A PDF that declares an encryption dictionary for the preflight to flag."""
    return build_pdf([["This synthetic document declares encryption."]], encrypt_marker=True)


QUALITY_REVIEW_PAGE_1 = (
    "Synthetic quality-review RFP QR-2026 - partial scan sample",
    "REQ-Q1: The service MUST include weekly status calls.",
    "REQ-Q2: The provider SHOULD deliver monthly reports.",
    "The supplier MUST provide encrypted backups.",
)


def build_quality_review_pdf() -> bytes:
    """A mixed PDF: one text page plus one deliberately blank (scan-like) page."""
    return build_pdf([QUALITY_REVIEW_PAGE_1, []])


def build_quality_review_workbook() -> Workbook:
    """Workbook with a modality-conflicting duplicate and an uncached formula."""
    workbook = Workbook()
    requirements = workbook.active
    if requirements is None:  # pragma: no cover - openpyxl always creates one sheet
        msg = "openpyxl workbook has no active sheet"
        raise RuntimeError(msg)
    requirements.title = "Requirements"
    requirements.append(["ID", "Requirement", "Modality"])
    requirements.append(["REQ-QX", "The provider SHOULD deliver monthly reports.", "MUST"])

    calc = workbook.create_sheet("Calc")
    calc.append(["Budget positions", 2])
    calc["C1"] = "=SUM(B1:B1)"

    workbook.properties.created = _FIXED_XLSX_TIMESTAMP
    workbook.properties.modified = _FIXED_XLSX_TIMESTAMP
    return workbook


def build_requirements_workbook() -> Workbook:
    """The synthetic requirements/timeline/formalities workbook."""
    workbook = Workbook()
    requirements = workbook.active
    if requirements is None:  # pragma: no cover - openpyxl always creates one sheet
        msg = "openpyxl workbook has no active sheet"
        raise RuntimeError(msg)
    requirements.title = "Requirements"
    requirements.append(["ID", "Requirement", "Modality"])
    requirements.append(
        ["REQ-X1", "The supplier MUST provide a German-speaking service desk.", "MUST"]
    )
    requirements.append(
        ["REQ-X2", "The service SHOULD provide quarterly service reviews.", "SHOULD"]
    )
    requirements.append(
        ["REQ-X3", "The supplier MUST encrypt customer data at rest and in transit.", "MUST"]
    )

    timeline = workbook.create_sheet("Timeline")
    timeline.append(["Milestone", "Date"])
    timeline.append(["Site visit", "2026-09-05 10:00"])
    timeline.append(["Supplier presentation", "2026-10-15 14:00"])

    formalities = workbook.create_sheet("Formalities")
    formalities.append(["Formality"])
    formalities.append(["All offer documents must be submitted in PDF format."])
    formalities.append(["The price sheet must be submitted as XLSX."])

    workbook.properties.created = _FIXED_XLSX_TIMESTAMP
    workbook.properties.modified = _FIXED_XLSX_TIMESTAMP
    return workbook


_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DOCX_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
    'officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_FIXED_ZIP_TIMESTAMP = (2026, 1, 1, 12, 0, 0)

#: SYNTHETIC RFP FIXTURE — invented for the requirement-to-golden-knowledge
#: bridge slice. No line is taken from any customer RFP. Uppercase
#: obligation keywords appear only in ``REQ-`` lines so the quality gate's
#: obligation coverage stays clean.
MTD_DOCX_PARAGRAPHS = (
    "SYNTHETIC RFP FIXTURE - Managed Detection RFP BF-2026-MTD-DEMO",
    "Customer: Example Testing Corp (synthetic test data)",
    "This request for proposal is fully synthetic and contains no real data.",
    "Submission deadline: 2026-11-30 17:00",
    "Offers must be submitted via the Example procurement portal.",
    "Proposals must be submitted in PDF format.",
    "Requirements",
    "REQ-M1: The provider MUST operate a 24x7 security monitoring service for the customer estate.",
    "REQ-M2: The service MUST support Microsoft Sentinel as the customer SIEM platform.",
    "REQ-M3: The provider MUST support Microsoft Defender XDR for endpoint detection.",
    "REQ-M4: The provider MUST be able to isolate compromised endpoints on request.",
    "REQ-M5: The provider SHOULD describe its transition and activation approach.",
    "REQ-M6: The provider MUST describe the handover to operations after go-live.",
    "REQ-M7: The platform MUST integrate with the QuantumTrace AlphaSIEM appliance.",
    "REQ-M8: Splunk support is not required and is excluded from scope.",
)


def build_mtd_requirements_docx() -> bytes:
    """A byte-deterministic minimal DOCX carrying the synthetic MTD RFP.

    Entries are stored uncompressed with fixed timestamps, so the bytes
    are reproducible independently of the zlib build.
    """
    body = "".join(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in MTD_DOCX_PARAGRAPHS)
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_WORD_NS}"><w:body>{body}</w:body></w:document>'
    ).encode()
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as archive:
        for name, payload in (
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES.encode("utf-8")),
            ("word/document.xml", document),
        ):
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_STORED
            archive.writestr(info, payload)
    return buffer.getvalue()


NOTES_TXT = (
    "Additional synthetic notes for the BidFusion walking skeleton.\n"
    "REQ-T1: The supplier MUST name a single point of contact.\n"
    "Late submissions must be submitted in writing with a justification.\n"
)
COVER_TXT = "Synthetic cover note for the blocked package fixture.\n"
REVIEW_NOTE_TXT = "Synthetic note accompanying the encrypted fixture.\n"


def write_fixtures(output: Path) -> list[Path]:
    """Write every synthetic fixture package below ``output``."""
    written: list[Path] = []

    def write(path: Path, content: bytes) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        written.append(path)

    minimal = output / "rfp-minimal"
    write(minimal / "rfp-invitation.pdf", build_invitation_pdf())
    write(minimal / "notes.txt", NOTES_TXT.encode("utf-8"))
    workbook_path = minimal / "requirements-matrix.xlsx"
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    build_requirements_workbook().save(workbook_path)
    written.append(workbook_path)

    blocked = output / "rfp-blocked"
    write(
        blocked / "payload.js",
        b"// Synthetic blocked-extension fixture for BidFusion tests.\n"
        b"// Inert text content; exists only to trigger the extension block.\n",
    )
    write(blocked / "cover.txt", COVER_TXT.encode("utf-8"))

    review = output / "rfp-review"
    write(review / "declared-encrypted.pdf", build_encrypted_marker_pdf())
    write(review / "note.txt", REVIEW_NOTE_TXT.encode("utf-8"))

    unsupported = output / "rfp-unsupported"
    write(unsupported / "blank-scan.pdf", build_blank_scan_pdf())

    mtd = output / "rfp-mtd"
    write(mtd / "mtd-requirements.docx", build_mtd_requirements_docx())

    quality = output / "rfp-quality-review"
    write(quality / "mixed-scan.pdf", build_quality_review_pdf())
    quality_workbook_path = quality / "quality-matrix.xlsx"
    quality_workbook_path.parent.mkdir(parents=True, exist_ok=True)
    build_quality_review_workbook().save(quality_workbook_path)
    written.append(quality_workbook_path)

    return written


def main() -> int:
    """Generate all fixture packages and report the written files."""
    parser = argparse.ArgumentParser(description="Generate synthetic BidFusion RFP fixtures.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    for path in write_fixtures(arguments.output):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
