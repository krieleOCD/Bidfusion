"""Self-tests for the technical test-isolation guards (Z5/H3).

These tests prove the guards themselves: network (TCP and UDP, however the
socket was obtained), subprocess execution, and direct reads of the
repository's protected ``inbox/``/``results/`` directories — including
descendants and symlinked routes — while ``tmp_path`` directories with the
same names, pypdf, and openpyxl keep working.
"""

from __future__ import annotations

import os
import shutil
import socket
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = REPO_ROOT / "fixtures" / "synthetic"


def test_tcp_and_udp_inet_sockets_cannot_be_created() -> None:
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.socket(socket.AF_INET6, socket.SOCK_DGRAM)
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.socket()  # the stdlib default family is AF_INET
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.create_connection(("192.0.2.1", 80), timeout=0.1)


def test_udp_sendto_and_connect_are_blocked_even_on_preexisting_sockets(
    unguarded_socket_class: type[socket.socket],
) -> None:
    """Even a socket created behind the factory guard cannot reach the net."""
    sock = unguarded_socket_class(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        with pytest.raises(RuntimeError, match="network access is disabled"):
            sock.sendto(b"x", ("192.0.2.1", 9))
        with pytest.raises(RuntimeError, match="network access is disabled"):
            sock.connect(("192.0.2.1", 9))
        with pytest.raises(RuntimeError, match="network access is disabled"):
            sock.connect_ex(("192.0.2.1", 9))
    finally:
        sock.close()


def test_local_ipc_stays_available() -> None:
    left, right = socket.socketpair()
    try:
        left.sendall(b"local ipc ok")
        assert right.recv(64) == b"local ipc ok"
    finally:
        left.close()
        right.close()


def test_dns_resolution_is_blocked() -> None:
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.getaddrinfo("example.com", 80)
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.gethostbyname("example.com")
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.gethostbyname_ex("example.com")
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.gethostbyaddr("192.0.2.1")
    with pytest.raises(RuntimeError, match="network access is disabled"):
        socket.getnameinfo(("192.0.2.1", 80), 0)


def test_subprocess_execution_is_blocked_in_regular_tests() -> None:
    with pytest.raises(RuntimeError, match="subprocess execution is disabled"):
        subprocess.run(["true"], check=False)
    with pytest.raises(RuntimeError, match="subprocess execution is disabled"):
        subprocess.Popen(["true"])
    with pytest.raises(RuntimeError, match="subprocess execution is disabled"):
        os.system("true")
    with pytest.raises(RuntimeError, match="subprocess execution is disabled"):
        os.popen("true")
    with pytest.raises(RuntimeError, match="subprocess execution is disabled"):
        os.posix_spawn("/usr/bin/true", ["true"], {})
    with pytest.raises(RuntimeError, match="subprocess execution is disabled"):
        os.fork()


@pytest.mark.allow_subprocess
def test_marked_tests_may_use_subprocess() -> None:
    completed = subprocess.run(["true"], check=False)
    assert completed.returncode == 0


@pytest.mark.parametrize("harness", ["inbox", "results"])
def test_direct_reads_of_protected_repository_paths_are_blocked(
    harness: str,
) -> None:
    target = REPO_ROOT / harness / "nonexistent.txt"
    with pytest.raises(PermissionError, match="blocked"):
        open(target, "rb")
    with pytest.raises(PermissionError, match="blocked"):
        target.read_bytes()
    with pytest.raises(PermissionError, match="blocked"):
        target.read_text()
    with pytest.raises(PermissionError, match="blocked"):
        os.open(target, os.O_RDONLY)


def test_protected_path_descendants_are_blocked() -> None:
    nested = REPO_ROOT / "inbox" / "sub" / "deeper" / "file.pdf"
    with pytest.raises(PermissionError, match="blocked"):
        nested.read_bytes()


def test_os_open_with_dir_fd_into_protected_paths_is_blocked() -> None:
    repo_fd = os.open(REPO_ROOT, os.O_RDONLY)
    try:
        for relative in ("inbox/secret.txt", "results", "results/run/x.json"):
            with pytest.raises(PermissionError, match="blocked"):
                os.open(relative, os.O_RDONLY, dir_fd=repo_fd)
    finally:
        os.close(repo_fd)


def test_os_open_with_dir_fd_outside_protected_paths_works(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "note.txt").write_text("synthetic", encoding="utf-8")
    base_fd = os.open(tmp_path, os.O_RDONLY)
    try:
        opened = os.open("inbox/note.txt", os.O_RDONLY, dir_fd=base_fd)
        os.close(opened)
    finally:
        os.close(base_fd)


def test_symlinked_routes_into_protected_paths_are_blocked(tmp_path: Path) -> None:
    link = tmp_path / "innocent"
    link.symlink_to(REPO_ROOT / "inbox")
    with pytest.raises(PermissionError, match="blocked"):
        (link / "file.txt").read_bytes()


def test_identically_named_tmp_directories_stay_usable(tmp_path: Path) -> None:
    for name in ("inbox", "results"):
        directory = tmp_path / name
        directory.mkdir()
        target = directory / "note.txt"
        target.write_text("synthetic", encoding="utf-8")
        assert target.read_text(encoding="utf-8") == "synthetic"
        descriptor = os.open(target, os.O_RDONLY)
        os.close(descriptor)


def test_pypdf_and_openpyxl_keep_working_over_tmp_path(tmp_path: Path) -> None:
    pdf_copy = tmp_path / "rfp-invitation.pdf"
    shutil.copyfile(FIXTURES / "rfp-minimal" / "rfp-invitation.pdf", pdf_copy)
    reader = PdfReader(BytesIO(pdf_copy.read_bytes()))
    assert len(reader.pages) == 2

    xlsx_copy = tmp_path / "requirements-matrix.xlsx"
    shutil.copyfile(FIXTURES / "rfp-minimal" / "requirements-matrix.xlsx", xlsx_copy)
    workbook = load_workbook(xlsx_copy, read_only=True, data_only=True)
    try:
        assert "Requirements" in workbook.sheetnames
    finally:
        workbook.close()
