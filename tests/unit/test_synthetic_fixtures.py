"""Guard tests for the committed synthetic fixture packages."""

from __future__ import annotations

import subprocess
import sys
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from bidfusion.schemas import ArtifactReference, ArtifactType, PreflightAssessment
from bidfusion.schemas.security import SecurityCheckStatus
from bidfusion.security import DeterministicPreflightScanner

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = REPO_ROOT / "fixtures" / "synthetic"
GENERATOR = REPO_ROOT / "scripts" / "generate_synthetic_fixtures.py"


def assess_fixture(path: Path) -> PreflightAssessment:
    content = path.read_bytes()
    artifact = ArtifactReference(
        artifact_id="ART-20260804-000001",
        bid_id="BID-2026-000123",
        run_id="RUN-20260804-0001",
        artifact_type=ArtifactType.RAW_DOCUMENT,
        schema_version="1.0.0",
        content_ref=f"package/{path.name}",
        content_hash=sha256(content).hexdigest(),
    )
    return DeterministicPreflightScanner().assess(
        artifact=artifact,
        document_id="DOC-001",
        filename=path.name,
        content=content,
    )


def check_status(assessment: PreflightAssessment, check: str) -> SecurityCheckStatus:
    return next(result.status for result in assessment.checks if result.check == check)


@pytest.mark.allow_subprocess
def test_pdf_fixtures_are_byte_reproducible(tmp_path: Path) -> None:
    regenerated = tmp_path / "regenerated"
    subprocess.run(
        [sys.executable, str(GENERATOR), "--output", str(regenerated)],
        check=True,
        capture_output=True,
    )
    for relative in (
        "rfp-minimal/rfp-invitation.pdf",
        "rfp-review/declared-encrypted.pdf",
        "rfp-unsupported/blank-scan.pdf",
        "rfp-quality-review/mixed-scan.pdf",
    ):
        assert (regenerated / relative).read_bytes() == (FIXTURES / relative).read_bytes(), relative


def test_invitation_pdf_has_two_readable_pages_with_expected_text() -> None:
    reader = PdfReader(BytesIO((FIXTURES / "rfp-minimal" / "rfp-invitation.pdf").read_bytes()))

    assert len(reader.pages) == 2
    page_1 = reader.pages[0].extract_text() or ""
    page_2 = reader.pages[1].extract_text() or ""
    assert "Submission deadline: 2026-09-30 17:00" in page_1
    assert "Q&A deadline: 2026-09-11 12:00" in page_1
    assert "The offer letter must be signed by an authorized representative." in page_1
    assert "REQ-P2: The supplier MUST encrypt customer data at rest and in transit." in page_2
    assert "REQ-P5: The solution handles peak load appropriately." in page_2


def test_invitation_pdf_page_provenance_is_stable_via_pypdf() -> None:
    reader = PdfReader(BytesIO((FIXTURES / "rfp-minimal" / "rfp-invitation.pdf").read_bytes()))
    texts = [page.extract_text() or "" for page in reader.pages]

    assert "REQ-P1" not in texts[0]
    assert "Submission deadline" not in texts[1]
    assert all(f"REQ-P{index}" in texts[1] for index in range(1, 6))


def test_blank_scan_pdf_is_structurally_valid_but_textless() -> None:
    content = (FIXTURES / "rfp-unsupported" / "blank-scan.pdf").read_bytes()
    reader = PdfReader(BytesIO(content))

    assert len(reader.pages) == 1
    assert "".join(page.extract_text() or "" for page in reader.pages).strip() == ""
    assert assess_fixture(FIXTURES / "rfp-unsupported" / "blank-scan.pdf").completed_successfully


def test_requirements_workbook_contains_expected_sheets_and_cells() -> None:
    workbook = load_workbook(
        BytesIO((FIXTURES / "rfp-minimal" / "requirements-matrix.xlsx").read_bytes()),
        read_only=True,
        data_only=True,
    )
    try:
        assert workbook.sheetnames == ["Requirements", "Timeline", "Formalities"]
        requirements = list(workbook["Requirements"].iter_rows())
        assert requirements[1][1].value == (
            "The supplier MUST provide a German-speaking service desk."
        )
        assert requirements[3][1].value == (
            "The supplier MUST encrypt customer data at rest and in transit."
        )
        timeline = list(workbook["Timeline"].iter_rows())
        assert timeline[1][1].value == "2026-09-05 10:00"
        assert timeline[2][1].value == "2026-10-15 14:00"
    finally:
        workbook.close()


def test_quality_review_package_carries_the_expected_review_triggers() -> None:
    pdf = (FIXTURES / "rfp-quality-review" / "mixed-scan.pdf").read_bytes()
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) == 2
    page_1 = reader.pages[0].extract_text() or ""
    assert "The supplier MUST provide encrypted backups." in page_1
    assert (reader.pages[1].extract_text() or "").strip() == ""

    workbook = load_workbook(
        BytesIO((FIXTURES / "rfp-quality-review" / "quality-matrix.xlsx").read_bytes()),
        read_only=True,
        data_only=True,
    )
    try:
        assert workbook.sheetnames == ["Requirements", "Calc"]
        calc = list(workbook["Calc"].iter_rows())
        assert len(calc[0]) < 3 or calc[0][2].value is None  # formula has no cached value
    finally:
        workbook.close()

    for name in ("mixed-scan.pdf", "quality-matrix.xlsx"):
        assessment = assess_fixture(FIXTURES / "rfp-quality-review" / name)
        assert assessment.completed_successfully, name


def test_clean_package_passes_the_existing_preflight() -> None:
    for path in sorted((FIXTURES / "rfp-minimal").iterdir()):
        assessment = assess_fixture(path)
        assert assessment.completed_successfully, path.name


def test_negative_fixtures_trigger_the_expected_checks() -> None:
    blocked = assess_fixture(FIXTURES / "rfp-blocked" / "payload.js")
    review = assess_fixture(FIXTURES / "rfp-review" / "declared-encrypted.pdf")

    assert check_status(blocked, "extension_mime") is SecurityCheckStatus.FAIL
    assert blocked.requires_gate_attention
    assert check_status(review, "encryption") is SecurityCheckStatus.FAIL
    assert review.requires_gate_attention
    assert assess_fixture(FIXTURES / "rfp-blocked" / "cover.txt").completed_successfully
    assert assess_fixture(FIXTURES / "rfp-review" / "note.txt").completed_successfully
