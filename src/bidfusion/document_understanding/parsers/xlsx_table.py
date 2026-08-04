"""Table extraction for regular XLSX workbooks via openpyxl.

Sheets become canonical pages (page number = sheet position), each with one
table whose cells keep sheet-qualified ``Sheet!A1`` references. Formulas are
never executed: with ``data_only=True`` openpyxl returns only stored cached
values, and formulas without a stored value are reported as warnings instead
of silently vanishing. Corrupt or unreadable workbooks fail closed.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from bidfusion.document_understanding.canonical import (
    DocumentParsingError,
    ParsedContent,
    UnsupportedDocumentError,
    block_source,
    verify_cleared_content,
)
from bidfusion.schemas.document import Page, Table, TableCell, TableRow
from bidfusion.schemas.ids import ArtifactId
from bidfusion.schemas.provenance import SourceLocator
from bidfusion.schemas.security import SecurityClearedArtifact

_CellKey = tuple[str, int, int]


class XlsxTableParser:
    """Extract sheet tables with cell-exact provenance from an XLSX workbook."""

    parser_name = "bidfusion-xlsx-table-parser"
    parser_version = "0.1.0"

    def parse(
        self,
        *,
        cleared: SecurityClearedArtifact,
        content: bytes,
        filename: str,
        canonical_artifact_id: ArtifactId,
    ) -> ParsedContent:
        """Parse cleared XLSX bytes into per-sheet canonical tables."""
        verify_cleared_content(cleared, content)
        try:
            sheet_values, formula_cells = _read_workbook(content)
        except Exception as failure:
            msg = "XLSX workbook could not be read"
            raise DocumentParsingError(msg) from failure

        pages: list[Page] = []
        warnings: list[str] = []
        stored_cells: set[_CellKey] = set()
        total_cells = 0
        for sheet_index, (sheet_title, cell_values) in enumerate(sheet_values, start=1):
            table_id = f"sheet-{sheet_index}"
            rows: dict[int, list[TableCell]] = {}
            for row_number, column_number, text in cell_values:
                coordinate = f"{sheet_title}!{get_column_letter(column_number)}{row_number}"
                stored_cells.add((sheet_title, row_number, column_number))
                rows.setdefault(row_number, []).append(
                    TableCell(
                        cell_id=coordinate,
                        text=text,
                        row=row_number,
                        column=column_number,
                        source=block_source(
                            cleared,
                            filename=filename,
                            canonical_artifact_id=canonical_artifact_id,
                            locator=SourceLocator(
                                page=sheet_index,
                                row=row_number,
                                column=column_number,
                                cell_id=coordinate,
                            ),
                            quote=text,
                        ),
                    )
                )

            if not rows:
                warnings.append(f"sheet {sheet_title!r} contains no cell content")
                pages.append(Page(page_number=sheet_index))
                continue

            table = Table(
                table_id=table_id,
                order=sheet_index - 1,
                caption=sheet_title,
                rows=tuple(
                    TableRow(
                        row_number=row_number,
                        cells=tuple(sorted(rows[row_number], key=lambda cell: cell.column)),
                    )
                    for row_number in sorted(rows)
                ),
                source=block_source(
                    cleared,
                    filename=filename,
                    canonical_artifact_id=canonical_artifact_id,
                    locator=SourceLocator(page=sheet_index, table_id=table_id),
                ),
            )
            total_cells += sum(len(row.cells) for row in table.rows)
            pages.append(Page(page_number=sheet_index, tables=(table,)))

        warnings.extend(
            f"formula without stored value skipped at {sheet}!{get_column_letter(column)}{row}"
            for sheet, row, column in sorted(formula_cells - stored_cells)
        )
        if total_cells == 0:
            msg = "XLSX workbook contains no readable cell content"
            raise UnsupportedDocumentError(msg)

        return ParsedContent(pages=tuple(pages), warnings=tuple(warnings))


def _read_workbook(
    content: bytes,
) -> tuple[list[tuple[str, list[tuple[int, int, str]]]], set[_CellKey]]:
    """Read stored values plus the coordinates of formula cells."""
    sheet_values: list[tuple[str, list[tuple[int, int, str]]]] = []
    values_workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        for worksheet in values_workbook.worksheets:
            cells: list[tuple[int, int, str]] = []
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None or cell.row is None or cell.column is None:
                        continue
                    text = str(value).strip()
                    if text:
                        cells.append((int(cell.row), int(cell.column), text))
            sheet_values.append((str(worksheet.title), cells))
    finally:
        values_workbook.close()

    formula_cells: set[_CellKey] = set()
    formulas_workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        for worksheet in formulas_workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not (isinstance(value, str) and value.startswith("=")):
                        continue
                    if cell.row is None or cell.column is None:
                        continue
                    formula_cells.add((str(worksheet.title), int(cell.row), int(cell.column)))
    finally:
        formulas_workbook.close()
    return sheet_values, formula_cells
